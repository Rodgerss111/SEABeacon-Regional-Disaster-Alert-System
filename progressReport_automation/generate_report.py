import os
import requests
import datetime
from dotenv import load_dotenv
from google import genai
from docx import Document

# Load environment variables
load_dotenv()
NOTION_API_KEY = os.getenv("NOTION_API_KEY")
NOTION_DATABASE_ID = os.getenv("NOTION_DATABASE_ID")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# Initialize Gemini Client
client = genai.Client(api_key=GEMINI_API_KEY)

def get_notion_tasks():
    """Fetches and filters tasks from Notion based on the current week."""
    url = f"https://api.notion.com/v1/databases/{NOTION_DATABASE_ID}/query"
    
    headers = {
        "Authorization": f"Bearer {NOTION_API_KEY}",
        "Notion-Version": "2022-06-28",
        "Content-Type": "application/json"
    }
    
    # Calculate the current week's start (Monday) and end (Sunday)
    today = datetime.date.today()
    start_of_week = today - datetime.timedelta(days=today.weekday())
    end_of_week = start_of_week + datetime.timedelta(days=6)
    
    print(f"Fetching tasks for the week of {start_of_week} to {end_of_week}...")
    
    response = requests.post(url, headers=headers)
    
    if response.status_code != 200:
        print(f"Error fetching from Notion: {response.text}")
        return []

    data = response.json()
    current_week_tasks = []
    
    for result in data.get("results", []):
        try:
            # 1. Extract Task Name (Updated to match your new "Task" column)
            task_name = result["properties"]["Task"]["title"][0]["plain_text"]
            
            # 2. Extract Status
            status = result["properties"]["Status"]["status"]["name"]
            
            if status not in ["Completed", "In Progress"]:
                continue

            # 3. Extract Dependencies (Capturing your detailed notes)
            try:
                # Assuming "Dependencies" is a standard Text property in Notion
                dependencies = result["properties"]["Dependencies"]["rich_text"][0]["plain_text"]
            except (KeyError, IndexError):
                dependencies = "No additional details."

            # 4. Extract Dates
            start_date_data = result["properties"]["Start Date"]["date"]
            end_date_data = result["properties"]["End Date"]["date"]
            
            start_date_str = start_date_data["start"] if start_date_data else None
            end_date_str = end_date_data["start"] if end_date_data else None
            
            if start_date_str and end_date_str:
                task_start = datetime.datetime.strptime(start_date_str, "%Y-%m-%d").date()
                task_end = datetime.datetime.strptime(end_date_str, "%Y-%m-%d").date()
                
                # Check if the task overlaps with the current week
                if task_start <= end_of_week and task_end >= start_of_week:
                    # Append everything nicely formatted for Gemini
                    current_week_tasks.append(
                        f"Task: {task_name}\nStatus: {status}\nDetails: {dependencies}\n---"
                    )
                    
        except (KeyError, IndexError, TypeError):
            continue
            
    return current_week_tasks

def generate_report_content(tasks):
    """Sends time-contextualized and detailed tasks to Gemini."""
    
    today = datetime.date.today()
    start_of_week = today - datetime.timedelta(days=today.weekday())
    end_of_week = start_of_week + datetime.timedelta(days=6)
    
    task_list_str = "\n".join(tasks)
    
    prompt = f"""
    Act as my thesis partner. I need to write 'Part I' and 'Part III' of my weekly progress report. 
    
    Context:
    - Today's date is {today.strftime('%B %d, %Y')}.
    - This report covers the week of {start_of_week.strftime('%B %d')} to {end_of_week.strftime('%B %d, %Y')}.
    
    Here are our active tasks for this specific week, including detailed notes from our Notion database:
    
    {task_list_str}
    
    Instructions:
    1. For Part I: Write a cohesive, professional paragraph summarizing our progress. Integrate the 'Details' provided for each task to make it sound descriptive and authentic. Write it in the first-person plural ('we'). Detail what was done day-by-day or task-by-task, note any challenges, and state the plan for next week based on what is 'In Progress'. Make it sound like a real human wrote it, not an AI, and use simple words.
    2. For Part III: Create a concise bulleted list of my specific individual workload based on these tasks. Make it sound like a real human wrote it, not an AI, and use simple words. Make each bullet point short and action-oriented, starting with a verb. Focus on what I personally did.
    
    Format your response EXACTLY like this:
    [PART 1 START]
    (Your paragraph here)
    [PART 1 END]
    
    [PART 3 START]
    (Your bullet points here)
    [PART 3 END]
    """
    
    print("Generating report via Gemini...")
    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt
    )
    
    text = response.text
    part_1 = text.split("[PART 1 START]")[1].split("[PART 1 END]")[0].strip()
    part_3 = text.split("[PART 3 START]")[1].split("[PART 3 END]")[0].strip()
    
    return part_1, part_3

def update_word_template(part_1_text, part_3_text):
    """Injects the Gemini text into the Word template."""
    # Load your master template (ensure this file exists in the same folder)
    template_path = "Master_Template.docx" 
    
    try:
        doc = Document(template_path)
    except Exception as e:
        print(f"Error: Could not find '{template_path}'. Please create it first.")
        return

    # Replace placeholders in the document
    # Ensure your Word doc has these exact strings on their own lines
    for paragraph in doc.paragraphs:
        if "[INSERT GEMINI PART I HERE]" in paragraph.text:
            paragraph.text = paragraph.text.replace("[INSERT GEMINI PART I HERE]", part_1_text)
        if "[INSERT GEMINI PART III HERE]" in paragraph.text:
            paragraph.text = paragraph.text.replace("[INSERT GEMINI PART III HERE]", part_3_text)

    # Save the new file with today's date
    today = datetime.datetime.now().strftime("%Y-%m-%d")
    output_filename = f"Navarro_Progress Report #{today}.docx"
    doc.save(output_filename)
    print(f"Success! Report saved as {output_filename}")

if __name__ == "__main__":
    print("Fetching tasks from Notion...")
    my_tasks = get_notion_tasks()
    
    if not my_tasks:
        print("No tasks found or error connecting to Notion.")
    else:
        print(f"Found {len(my_tasks)} tasks.")
        part1, part3 = generate_report_content(my_tasks)
        update_word_template(part1, part3)